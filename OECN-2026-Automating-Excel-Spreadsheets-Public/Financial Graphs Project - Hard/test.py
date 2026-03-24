

import pandas as pd
from pandas import DataFrame
from openpyxl import Workbook, worksheet, load_workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.styles import Color
from openpyxl.drawing.colors import ColorChoice
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker
from openpyxl.utils.dataframe import dataframe_to_rows


#Example of how to create a line chart
def example_line_chart(output_file:str):
    print("Example Chart")
    # Sample data
    data = [
        ["Month", "2023"],
        ["Jan", 10],
        ["Feb", 15],
        ["Mar", 12],
        ["Apr", 18],
        ["May", 20],
    ]

    wb = Workbook()
    ws = wb.active

    for row in data:
        ws.append(row)

    data_values = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=len(data))
    categories = Reference(ws, min_col=1, min_row=2, max_col=1, max_row=len(data))

    chart = LineChart()
    chart.title = "Example Line Chart"
    chart.style = 20
    chart.add_data(data_values)
    chart.set_categories(categories)

    ws.add_chart(chart, "D1")  # Place chart in Cell D1
    wb.save(output_file)
    print("example.xlsx created. Open it to see the chart.")

if __name__ == "__main__":
    output_file = "example.xlsx"
    example_line_chart(output_file)  